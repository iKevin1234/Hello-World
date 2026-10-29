"""Read-only Cisco access VLAN audit. Run with --help for options."""
import argparse
from collections import Counter
from datetime import datetime
from getpass import getpass
from pathlib import Path
import re
import sys


def parse_switchports(output):
    """Count configured access mode, never infer it from operational mode."""
    if re.search(r"(?im)^\s*(%\s*(invalid|error|incomplete|ambiguous|authorization)|error:)", output):
        raise ValueError("Switch rejected the show command")
    blocks = re.split(r"(?im)^\s*Name:\s*", output)[1:]
    if not blocks:
        raise ValueError("No switchport records recognized; check platform/command permissions")
    ports = {}
    for block in blocks:
        name = block.splitlines()[0].strip()
        # Physical Ethernet ports only; excludes port-channels and subinterfaces.
        if not re.fullmatch(r"[A-Za-z-]+\d+(?:/\d+)+", name):
            continue
        fields = {}
        for line in block.splitlines()[1:]:
            key, sep, value = line.partition(":")
            if sep:
                fields[key.strip().lower()] = value.strip()
        if fields.get("switchport", "").lower() == "disabled":
            continue
        mode = fields.get("administrative mode", "").lower()
        if not mode:
            raise ValueError(f"Missing administrative mode on {name}")
        if mode not in {"static access", "access"}:
            continue
        match = re.match(r"(\d+)\b", fields.get("access mode vlan", ""))
        if not match or not 1 <= int(match[1]) <= 4094:
            raise ValueError(f"Missing/unsupported access VLAN on {name}")
        if name in ports:
            raise ValueError(f"Duplicate interface record: {name}")
        ports[name] = int(match[1])
    return ports


def interface_key(name):
    match = re.fullmatch(r"([A-Za-z-]+)(\d+(?:/\d+)+)", name.strip())
    if not match:
        return name.lower()
    prefix = match[1].lower()
    aliases = {"gigabitethernet": "gi", "fastethernet": "fa", "tengigabitethernet": "te",
               "twentyfivegige": "twe", "fortygigabitethernet": "fo", "hundredgige": "hu",
               "ethernet": "eth", "et": "eth", "two-gigabitethernet": "tw"}
    return aliases.get(prefix, prefix) + match[2]


def parse_link_status(output):
    """Read the Status column, avoiding interface descriptions containing status words."""
    states = {}
    start = end = None
    down_states = {"notconnect", "disabled", "err-disabled", "inactive", "sfpabsent",
                   "xcvrabsent", "suspended", "down", "nooper", "linkfault", "channeldo"}
    for line in output.splitlines():
        header = re.search(r"\bStatus\s+Vlan\b", line, re.I)
        if header and re.match(r"\s*Port\b", line, re.I):
            start = header.start()
            end = re.search(r"\bVlan\b", line, re.I).start()
            continue
        if start is None:
            continue
        port = re.match(r"\s*([A-Za-z-]+\d+(?:/\d+)+)\s+", line)
        if not port:
            continue
        state = line[start:end].strip().lower()
        if state in {"connected", "up"}:
            states[interface_key(port[1])] = "up"
        elif state in down_states:
            states[interface_key(port[1])] = "down"
    if not states:
        raise ValueError("No link status recognized from show interfaces status")
    return states


def count_vlan_links(ports, states):
    totals, up, down = Counter(ports.values()), Counter(), Counter()
    for name, vlan in ports.items():
        state = states.get(interface_key(name))
        if state not in {"up", "down"}:
            raise ValueError(f"Missing or unrecognized link status for {name}; cannot report accurate up/down counts")
        (up if state == "up" else down)[vlan] += 1
    return [(vlan, count, up[vlan], down[vlan]) for vlan, count in sorted(totals.items())]


def read_inventory(path, sheet=None):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Install dependencies: python -m pip install netmiko openpyxl") from exc
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = book[sheet] if sheet else book.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        headers = [re.sub(r"[^a-z0-9]", "", str(v or "").lower()) for v in next(rows, ())]
        def column(*names):
            return next((headers.index(n) for n in names if n in headers), None)
        host_col = column("ipaddress", "ip", "managementip", "address", "device", "host")
        name_col = column("hostname", "name", "switch")
        type_col = column("devicetype", "platform")
        if host_col is None:
            host_col = name_col
        if host_col is None:
            raise ValueError("Excel needs an IP Address, Device, Host, or Hostname header in row 1")
        devices, seen = [], set()
        for row_number, row in enumerate(rows, 2):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            def value(index):
                return str(row[index]).strip() if index is not None and row[index] is not None else ""
            host = value(host_col)
            if not host:
                raise ValueError(f"Inventory row {row_number}: missing device address")
            if host.lower() in seen:
                raise ValueError(f"Inventory row {row_number}: duplicate address {host}")
            seen.add(host.lower())
            kind = value(type_col).lower() or "cisco_ios"
            kind = {"ios": "cisco_ios", "ios-xe": "cisco_ios", "iosxe": "cisco_ios", "nx-os": "cisco_nxos", "nxos": "cisco_nxos"}.get(kind, kind)
            if kind not in {"cisco_ios", "cisco_xe", "cisco_nxos"}:
                raise ValueError(f"Inventory row {row_number}: unsupported device type {kind}")
            devices.append({"host": host, "hostname": value(name_col) or host, "device_type": kind})
        if not devices:
            raise ValueError("Inventory contains no devices")
        return devices
    finally:
        book.close()


def parse_stack(output):
    """Read Catalyst show switch; exclude absent/provisioned member slots."""
    rows = re.findall(
        r"(?im)^\s*\*?\s*(\d+)\s+(?:Active|Standby|Master|Member|Slave)\s+"
        r"([0-9a-f]{4}\.[0-9a-f]{4}\.[0-9a-f]{4})\s+\d+\s+(.*?)\s*$", output)
    if not rows:
        return "Unknown", "", "show switch unsupported, denied, or unrecognized", set()
    present, absent, warnings = set(), set(), []
    for number, mac, remainder in rows:
        member = int(number)
        if mac == "0000.0000.0000" or re.search(r"\b(?:Provisioned|Unprovisioned|Removed|Absent)\b", remainder, re.I):
            absent.add(member)
        else:
            present.add(member)
            if not re.search(r"\bReady\s*$", remainder, re.I):
                warnings.append(f"Member {member}: {remainder}")
    if absent:
        warnings.append("Excluded absent/provisioned slots: " + ", ".join(map(str, sorted(absent))))
    if not present:
        return "Unknown", "", "; ".join(warnings) or "No present members identified", absent
    return "Yes" if len(present) > 1 else "No", len(present), "; ".join(warnings), absent


def collect_stack(connection, device_type, timeout):
    if device_type == "cisco_nxos":
        return "Unknown", "", "Catalyst stack detection not applicable to NX-OS", set()
    try:
        return parse_stack(connection.send_command("show switch", read_timeout=timeout))
    except Exception as exc:
        # A failed optional check must not discard already-collected VLAN data.
        return "Unknown", "", f"Stack check failed: {type(exc).__name__}", set()


def write_report(path, summary, status):
    """Save a formatted workbook atomically, preserving the previous checkpoint."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    sheets = [
        ("VLAN Summary", ["Hostname", "IP Address", "Access VLAN", "Total Configured Access Ports", "Up Ports", "Down Ports"], summary, [32, 22, 16, 24, 16, 16]),
        ("Device Status", ["Hostname", "IP Address", "Status", "Total Access Port Count", "Credential Set", "Error", "Stacked Switch", "Stack Members", "Stack Notes"], status, [32, 22, 14, 26, 20, 80, 20, 20, 80]),
    ]
    book = Workbook()
    book.remove(book.active)
    edge = Side(style="thin", color="B7C9D6")
    cell_border = Border(left=edge, right=edge, top=edge, bottom=edge)
    for name, headers, rows, widths in sheets:
        if len(rows) > 1_048_575:
            raise ValueError(f"{name} exceeds Excel's worksheet row limit")
        sheet = book.create_sheet(name)
        sheet.append(headers)
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 28
        for cell in sheet[1]:
            cell.border = cell_border
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        previous_device = None
        for row in rows:
            display_row = list(row)
            if name == "VLAN Summary":
                device = tuple(row[:2])
                if device == previous_device:
                    display_row[:2] = [None, None]
                previous_device = device
            sheet.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in display_row])
            for cell in sheet[sheet.max_row]:
                cell.border = cell_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                # Preserve device text literally, including strings starting with '='.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                elif isinstance(cell.value, int):
                    cell.number_format = "#,##0"
            if name == "VLAN Summary" and display_row[2] == "0 access ports configured":
                sheet.cell(sheet.max_row, 3).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                sheet.row_dimensions[sheet.max_row].height = 45
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.auto_filter.ref = sheet.dimensions
        if name == "Device Status":
            for row in sheet.iter_rows(min_row=2):
                row[5].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                row[8].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if row[8].value:
                    sheet.row_dimensions[row[0].row].height = 60
                if row[2].value == "ERROR":
                    row[2].fill = PatternFill("solid", fgColor="FCE4D6")
                    sheet.row_dimensions[row[0].row].height = 60
    temporary = path.with_name(path.stem + ".tmp.xlsx")
    try:
        book.save(temporary)
        temporary.replace(path)
    except PermissionError as exc:
        raise PermissionError(f"Cannot update {path}. Close it in Excel before running the script. Latest workbook may be available at {temporary}.") from exc
    finally:
        book.close()


def connect_with_fallback(device, tacacs, local, secret=""):
    """Retry only SSH authentication rejection; AAA policy stays on the switch."""
    from netmiko import ConnectHandler
    from netmiko.exceptions import NetmikoAuthenticationException

    options = dict(host=device["host"], device_type=device["device_type"],
                   secret=secret, conn_timeout=20, auth_timeout=30,
                   banner_timeout=30, use_keys=False, allow_agent=False)
    try:
        return ConnectHandler(**options, **tacacs), "TACACS"
    except NetmikoAuthenticationException:
        print("  TACACS credentials rejected; trying local credentials", flush=True)
    return ConnectHandler(**options, **local), "Local"


def hostname_from_prompt(prompt):
    """Extract a hostname from a normal Cisco EXEC prompt."""
    match = re.fullmatch(r"([A-Za-z0-9_][A-Za-z0-9_.-]*)(?:\(standby\))?[#>]", prompt.strip())
    return match[1] if match else None


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inventory", type=Path, nargs="?", help="Defaults to devices.xlsx beside this script; reads the first worksheet")
    parser.add_argument("--host", help="Test one switch directly instead of using an Excel inventory")
    parser.add_argument("--device-type", choices=["cisco_ios", "cisco_xe", "cisco_nxos"], default="cisco_ios", help="Platform for --host")
    parser.add_argument("--sheet", help="Inventory worksheet name")
    parser.add_argument("--username", help="TACACS username; prompted if omitted")
    parser.add_argument("--local-username", help="Fallback local username; prompted if omitted")
    parser.add_argument("--enable", action="store_true", help="Prompt for a separate enable secret")
    parser.add_argument("--output-dir", type=Path, default=Path("vlan_reports"))
    parser.add_argument("--timeout", type=int, default=120, help="Show command timeout in seconds")
    args = parser.parse_args(argv)
    if args.timeout <= 0:
        parser.error("--timeout must be positive")
    if args.inventory and args.host:
        parser.error("Provide either an Excel inventory or --host, but not both")
    if not args.host:
        args.inventory = args.inventory or Path(__file__).resolve().parent / "devices.xlsx"
        if not args.inventory.is_file():
            parser.error(f"Inventory not found: {args.inventory}. Place devices.xlsx in the same folder as this script.")
        print(f"Using inventory: {args.inventory.resolve()}", flush=True)
    try:
        import openpyxl  # Required for output even when using --host.
        devices = ([{"host": args.host, "hostname": args.host, "device_type": args.device_type}]
                   if args.host else read_inventory(args.inventory, args.sheet))
        from netmiko import ConnectHandler
    except ImportError:
        parser.error("Install dependencies with this Python: python -m pip install netmiko openpyxl")
    except (ValueError, OSError, KeyError) as exc:
        parser.error(str(exc))
    username = args.username or input("TACACS username: ").strip()
    password = getpass("TACACS password: ")
    local_username = args.local_username or input("Local fallback username: ").strip()
    local_password = getpass("Local fallback password: ")
    if not username or not local_username:
        parser.error("Both TACACS and local usernames are required")
    tacacs = {"username": username, "password": password}
    local = {"username": local_username, "password": local_password}
    secret = getpass("Enable secret: ") if args.enable else ""
    run_dir = args.output_dir / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    run_dir.mkdir(parents=True, exist_ok=False)
    report_path = run_dir / "access_vlan_report.xlsx"
    summary, status = [], []
    for index, device in enumerate(devices, 1):
        host, hostname = device["host"], device["hostname"]
        print(f"[{index}/{len(devices)}] {hostname} ({host})", flush=True)
        credential_set = "None"
        stacked, members, stack_notes, absent_members = "Unknown", "", "Not collected", set()
        try:
            connection, credential_set = connect_with_fallback(device, tacacs, local, secret)
            with connection:
                if args.enable:
                    connection.enable()
                discovered_hostname = hostname_from_prompt(connection.find_prompt())
                if discovered_hostname:
                    hostname = discovered_hostname
                    print(f"  Device hostname: {hostname}", flush=True)
                else:
                    print(f"  Hostname not recognized in CLI prompt; using inventory/address: {hostname}", flush=True)
                command = "show interface switchport" if device["device_type"] == "cisco_nxos" else "show interfaces switchport"
                output = connection.send_command(command, read_timeout=args.timeout)
                status_command = "show interface status" if device["device_type"] == "cisco_nxos" else "show interfaces status"
                link_output = connection.send_command(status_command, read_timeout=args.timeout)
                stacked, members, stack_notes, absent_members = collect_stack(connection, device["device_type"], args.timeout)
            ports = parse_switchports(output)
            # Stack interfaces use member/slot/port. Do not count ghost ports on absent slots.
            ports = {name: vlan for name, vlan in ports.items()
                     if not ((match := re.fullmatch(r"[A-Za-z-]+(\d+)/\d+/\d+", name))
                             and int(match[1]) in absent_members)}
            counts = count_vlan_links(ports, parse_link_status(link_output)) if ports else []
            for vlan, count, up, down in counts:
                summary.append([hostname, host, vlan, count, up, down])
                print(f"  VLAN {vlan}: {count} ports ({up} up, {down} down)")
            status.append([hostname, host, "OK", len(ports), credential_set, "", stacked, members, stack_notes])
            if not ports:
                summary.append([hostname, host, "0 access ports configured", 0, 0, 0])
                print("  No configured physical access ports")
        except Exception as exc:
            message = str(exc)
            for credential in sorted({password, local_password, secret}, key=len, reverse=True):
                if credential:
                    message = message.replace(credential, "[redacted]")
            status.append([hostname, host, "ERROR", "", credential_set, f"{type(exc).__name__}: {message}", stacked, members, stack_notes])
            print(f"  ERROR: {type(exc).__name__}; see the Device Status tab")
        # Checkpoint after each device so completed results survive later interruptions.
        write_report(report_path, summary, status)
    print(f"Excel report saved to: {report_path.resolve()}")
    return 1 if any(row[2] == "ERROR" for row in status) else 0


if __name__ == "__main__":
    sys.exit(main())
